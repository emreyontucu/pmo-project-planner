import datetime
import io
import re
import urllib.parse
from typing import List

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from ..audit_service import log_audit
from ..database import get_db
from ..models import Project, ProjectAuditLog
from ..scheduling import scheduling_engine
from .. import schemas

router = APIRouter(prefix="/projects", tags=["projects"])


async def _get_project_or_404(project_id: int, db: AsyncSession) -> Project:
    result = await db.execute(
        select(Project)
        .options(
            selectinload(Project.tasks),
            selectinload(Project.dependencies),
            selectinload(Project.answers),
            selectinload(Project.audit_logs),
        )
        .where(Project.id == project_id)
    )
    project = result.scalar_one_or_none()
    if project is None:
        raise HTTPException(status_code=404, detail="Proje bulunamadı.")
    return project


@router.post("", response_model=schemas.ProjectRead, status_code=201)
async def create_project(payload: schemas.ProjectCreate, db: AsyncSession = Depends(get_db)):
    existing = await db.execute(select(Project).where(Project.code == payload.code))
    if existing.scalar_one_or_none() is not None:
        raise HTTPException(status_code=400, detail=f"'{payload.code}' kodlu proje zaten mevcut.")

    project = Project(**payload.model_dump())
    db.add(project)
    await db.flush()
    await db.refresh(project)
    await log_audit(db, project.id, "PROJE_OLUŞTURULDU", f"'{project.name}' ({project.code}) projesi oluşturuldu.")
    return project


@router.get("", response_model=List[schemas.ProjectRead])
async def list_projects(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Project).order_by(Project.id))
    return result.scalars().all()


@router.get("/{project_id}", response_model=schemas.ProjectDetail)
async def get_project(project_id: int, db: AsyncSession = Depends(get_db)):
    return await _get_project_or_404(project_id, db)


@router.get("/{project_id}/audit-logs", response_model=List[schemas.ProjectAuditLogRead])
async def get_project_audit_logs(project_id: int, db: AsyncSession = Depends(get_db)):
    await _get_project_or_404(project_id, db)
    result = await db.execute(
        select(ProjectAuditLog)
        .where(ProjectAuditLog.project_id == project_id)
        .order_by(ProjectAuditLog.created_at.desc(), ProjectAuditLog.id.desc())
    )
    return result.scalars().all()


@router.patch("/{project_id}", response_model=schemas.ProjectRead)
async def update_project(project_id: int, payload: schemas.ProjectUpdate, db: AsyncSession = Depends(get_db)):
    project = await _get_project_or_404(project_id, db)
    updates = payload.model_dump(exclude_unset=True)
    changed_fields = []
    for field, value in updates.items():
        setattr(project, field, value)
        changed_fields.append(f"{field}: {value}")
    await log_audit(db, project.id, "PROJE_GÜNCELLENDİ", f"Proje bilgileri güncellendi ({', '.join(changed_fields)}).")
    await db.flush()
    await db.refresh(project)
    return project


@router.patch("/{project_id}/cancel", response_model=schemas.ProjectRead)
async def cancel_project(project_id: int, db: AsyncSession = Depends(get_db)):
    project = await _get_project_or_404(project_id, db)
    project.business_status = "İptal Edildi"
    await log_audit(db, project.id, "PROJE_İPTAL_EDİLDİ", f"'{project.name}' projesinin durumu 'İptal Edildi' olarak güncellendi.")
    await db.flush()
    await db.refresh(project)
    return project


@router.get("/{project_id}/export")
async def export_project_tasks(project_id: int, db: AsyncSession = Depends(get_db)):
    """Exports all tasks of the project in the standard Görev Takip Excel format."""
    project = await _get_project_or_404(project_id, db)
    
    def sort_key(t):
        phase_str = t.phase or ""
        match = re.match(r"^[sS]([1-5])", phase_str.strip())
        phase_num = int(match.group(1)) if match else 99
        return (phase_num, t.order_index, t.start_date or datetime.date(2099, 1, 1), t.id)

    sorted_tasks = sorted(project.tasks, key=sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = "Proje Planı"

    headers = [
        "Proje Aşaması",
        "Görev Adı",
        "Sorumlu",
        "Öncelik",
        "Durum",
        "Planlanan Başlangıç Tarihi",
        "Gerçekleşen Başlangıç Tarihi",
        "Planlanan Bitiş Tarihi",
        "Gerçekleşen Bitiş Tarihi",
        "Notlar",
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="365F91", end_color="365F91", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_side = Side(style="thin", color="D9D9D9")
    row_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    ws.row_dimensions[1].height = 28
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    def format_date_str(d: datetime.date | None) -> str:
        return d.strftime("%d.%m.%Y") if d else ""

    for t in sorted_tasks:
        row_data = [
            t.phase or "",
            t.name,
            t.sorumlu or "",
            t.priority or "",
            t.status or "Başlamadı",
            format_date_str(t.start_date),
            format_date_str(t.actual_start_date),
            format_date_str(t.end_date),
            format_date_str(t.actual_end_date),
            t.description or "",
        ]
        ws.append(row_data)

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    for row_idx in range(2, len(sorted_tasks) + 2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Calibri", size=10)
            cell.border = row_border
            if col_idx in (4, 5, 6, 7, 8, 9):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    clean_name = re.sub(r'[\/:*?"<>|]', "_", project.name)
    encoded_filename = urllib.parse.quote(f"{clean_name}_Proje_Plani.xlsx")
    headers_resp = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
        "Access-Control-Expose-Headers": "Content-Disposition",
    }
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_resp,
    )


@router.delete("/{project_id}", status_code=204)
async def delete_project(project_id: int, db: AsyncSession = Depends(get_db)):
    project = await _get_project_or_404(project_id, db)
    await db.delete(project)
    return None


@router.post("/{project_id}/schedule", response_model=schemas.ProjectDetail)
async def schedule_project(project_id: int, db: AsyncSession = Depends(get_db)):
    """Recalculates start/end dates for all tasks using the dependency graph and company calendar."""
    project = await _get_project_or_404(project_id, db)
    if project.start_date is None:
        raise HTTPException(status_code=400, detail="Tarih hesaplaması için projenin başlangıç tarihi gerekli.")

    try:
        scheduling_engine.calculate_dates(
            project.start_date,
            project.tasks,
            project.dependencies,
            exclude_bridge_days=project.exclude_bridge_days,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    await log_audit(db, project.id, "TARİHLER_HESAPLANDI", "Bağımlılık grafiği ve çalışma takvimine göre tüm görev tarihleri otomatik yeniden hesaplandı.")
    await db.flush()
    await db.refresh(project)
    return project


@router.post("/{project_id}/finalize", response_model=schemas.FinalizeResult)
async def finalize_project(project_id: int, db: AsyncSession = Depends(get_db)):
    """Publishes the plan as Final. Blocked while any task is missing a start/end date."""
    project = await _get_project_or_404(project_id, db)

    missing = [t.id for t in project.tasks if t.start_date is None or t.end_date is None]
    if missing:
        raise HTTPException(
            status_code=400,
            detail={
                "message": f"Planınızda henüz tarih atanmamış {len(missing)} task bulunuyor.",
                "missing_dates_task_ids": missing,
            },
        )

    project.status = "Final"
    await log_audit(db, project.id, "PLAN_YAYINLANDI", "Proje planı 'Final' olarak onaylandı ve planlanan tarihler kilitlendi.")
    await db.flush()
    return schemas.FinalizeResult(status=project.status, missing_dates_task_ids=[])
