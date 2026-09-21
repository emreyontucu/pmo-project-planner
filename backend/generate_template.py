import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()

# Main sheet: Görev Takip
ws = wb.active
ws.title = "Görev Takip"

headers = [
    "Sıra No",
    "Proje Aşaması",
    "Görev Adı",
    "Öncelik",
    "Sorumlu",
    "Planlanan Başlangıç",
    "Planlanan Bitiş",
    "Gerçekleşen Başlangıç",
    "Gerçekleşen Bitiş",
    "Durum",
    "Not",
]
ws.append(headers)

# Styling
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
border_side = Side(style="thin", color="CBD5E1")
row_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

ws.row_dimensions[1].height = 30
for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align

# Sample tasks to guide user
sample_data = [
    [1, "S1 - İş ve Veri Analizi / Kavramsal Tasarım", "Gereksinim ve Kapsam Analizi", "Yüksek", "İş Analisti", "07.09.2026", "11.09.2026", "", "", "Başlamadı", "İş birimi paydaş mülakatları"],
    [2, "S1 - İş ve Veri Analizi / Kavramsal Tasarım", "Teknik Mimari & API Sözleşmesi", "Yüksek", "Çözüm Mimarı", "14.09.2026", "18.09.2026", "", "", "Başlamadı", "Servis spesifikasyonları"],
    [3, "S2 - Veri Toplama ve Hazırlık", "Veri Haritalama ve Şema Tasarımı", "Orta", "Veri Mühendisi", "21.09.2026", "25.09.2026", "", "", "Başlamadı", "Kaynak ve hedef tablolar"],
    [4, "S3 - Geliştirme / Modelleme", "Temel Modül ve Arayüz Geliştirmesi", "Kritik", "Yazılım Ekibi", "28.09.2026", "09.10.2026", "", "", "Başlamadı", "Sprint 1 geliştirme"],
    [5, "S4 - Test / Model Değerlendirme", "Fonksiyonel ve Entegrasyon Testleri", "Yüksek", "QA Test Uzmanı", "12.10.2026", "16.10.2026", "", "", "Başlamadı", "Uçtan uca test senaryoları"],
    [6, "S5 - Canlı Geçiş / Gözlem / Kapanış", "Kullanıcı Eğitimi ve Canlıya Geçiş", "Yüksek", "Proje Yöneticisi", "19.10.2026", "23.10.2026", "", "", "Başlamadı", "Pilot devreye alım"],
]

center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")

for row_idx, row_values in enumerate(sample_data, start=2):
    ws.row_dimensions[row_idx].height = 22
    for col_idx, val in enumerate(row_values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = Font(name="Calibri", size=10)
        cell.border = row_border
        if col_idx in (1, 4, 6, 7, 8, 9, 10):
            cell.alignment = center_align
        else:
            cell.alignment = left_align

# Adjust column widths
for col in ws.columns:
    col_letter = get_column_letter(col[0].column)
    max_len = max(len(str(cell.value or "")) for cell in col)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

# Set widths specifically for readability
ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 38
ws.column_dimensions["C"].width = 40
ws.column_dimensions["D"].width = 14
ws.column_dimensions["E"].width = 24
ws.column_dimensions["F"].width = 20
ws.column_dimensions["G"].width = 20
ws.column_dimensions["H"].width = 20
ws.column_dimensions["I"].width = 20
ws.column_dimensions["J"].width = 16
ws.column_dimensions["K"].width = 35

# Add Data Validations for Öncelik and Durum
dv_priority = DataValidation(type="list", formula1='"Kritik,Yüksek,Orta,Düşük"', allow_blank=True)
ws.add_data_validation(dv_priority)
dv_priority.add("D2:D500")

dv_status = DataValidation(type="list", formula1='"Başlamadı,Devam Ediyor,Beklemede,Tamamlandı,İptal Edildi"', allow_blank=True)
ws.add_data_validation(dv_status)
dv_status.add("J2:J500")

dv_phase = DataValidation(type="list", formula1='"S1 - İş ve Veri Analizi / Kavramsal Tasarım,S2 - Veri Toplama ve Hazırlık,S3 - Geliştirme / Modelleme,S4 - Test / Model Değerlendirme,S5 - Canlı Geçiş / Gözlem / Kapanış"', allow_blank=True)
ws.add_data_validation(dv_phase)
dv_phase.add("B2:B500")

wb.save("/Users/emre/pmo-project-planner/frontend/public/Gorev_Takip_Sablonu.xlsx")
print("New template successfully generated and saved to frontend/public/Gorev_Takip_Sablonu.xlsx")
